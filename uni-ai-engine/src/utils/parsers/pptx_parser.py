import io
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from utils.parsers.base import BaseParser
from utils.loggers.app_logger import app_logger


class PptxParser(BaseParser):
    def extract_text(self, file_bytes: bytes) -> str:
        try:
            stream = io.BytesIO(file_bytes)
            prs = Presentation(stream)
            text_parts = []

            processed_image_hashes = set()

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = []
                image_texts = []

                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())

                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_bytes = shape.image.blob

                        img_hash = hash(image_bytes)
                        if img_hash not in processed_image_hashes:
                            ocr_result = self.ocr_image_bytes(image_bytes)
                            if ocr_result:
                                image_texts.append(
                                    f"[Nội dung OCR từ ảnh]: {ocr_result}")
                            processed_image_hashes.add(img_hash)

                if slide_text or image_texts:
                    text_parts.append(f"--- Nội dung Slide {slide_num} ---")
                    if slide_text:
                        text_parts.extend(slide_text)
                    if image_texts:
                        text_parts.extend(image_texts)
                    text_parts.append("\n")

            return "\n".join(text_parts)
        except Exception as e:
            app_logger.error(f"[Error] Lỗi đọc PPTX: {e}")
            return ""
