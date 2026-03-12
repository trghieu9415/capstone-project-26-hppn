import io
from docx import Document
from utils.logger import app_logger
from utils.parsers.base import BaseParser


class DocxParser(BaseParser):
    def extract_text(self, file_bytes: bytes) -> str:
        try:
            stream = io.BytesIO(file_bytes)
            doc = Document(stream)
            text_parts = []

            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text.strip())

            for table in doc.tables:
                text_parts.append("\n[Bắt đầu bảng]")
                for row in table.rows:
                    row_data = [cell.text.replace("\n", " ").strip() for cell in
                                row.cells]
                    if any(row_data):
                        text_parts.append(" | ".join(row_data))
                text_parts.append("[Kết thúc bảng]\n")

            processed_image_parts = set()
            image_count = 0

            for rel in doc.part.rels.values():
                if "image" in rel.target_ref and rel.target_part not in processed_image_parts:
                    image_bytes = rel.target_part.blob
                    ocr_text = self.ocr_image_bytes(image_bytes)

                    if ocr_text:
                        image_count += 1
                        text_parts.append(
                            f"\n[Nội dung OCR từ ảnh {image_count}]:\n{ocr_text}\n")

                    processed_image_parts.add(rel.target_part)

            return "\n".join(text_parts)
        except Exception as e:
            app_logger.error(f"[Error] Lỗi đọc DOCX: {e}")
            return ""
